import tkinter as tk
from tkinter import filedialog, messagebox
import json
import os
import sys
import csv
from openpyxl import load_workbook, Workbook

CONFIG_FILE = os.path.expanduser("~/.fai_mapping_config.json")

def load_config():
    if os.path.exists(CONFIG_FILE):
        try:
            with open(CONFIG_FILE, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            pass
    return {"template_path":"", "title_path":"", "output_path":""}

def save_config(cfg):
    with open(CONFIG_FILE, "w", encoding="utf-8") as f:
        json.dump(cfg, f, ensure_ascii=False, indent=2)

def col_num_to_letter(n):
    """数字列号转Excel列字母，1->A,26->Z,27->AA"""
    if n is None:
        return None
    s = ""
    while n > 0:
        n, rem = divmod(n - 1, 26)
        s = chr(ord('A') + rem) + s
    return s

def run_mapping(template_path, title_path, output_path):
    # ========= 1.读取Template CPK工作表，全部行一次性读入内存 =========
    wb_temp = load_workbook(template_path, read_only=True, data_only=True)
    ws_cpk = wb_temp["CPK"]
    all_rows = list(ws_cpk.iter_rows(values_only=True))
    wb_temp.close()

    # 扫描前20行找表头FAI#ROW
    fai_col_idx = None
    start_data_row_idx = None
    for scan_idx, row_tuple in enumerate(all_rows[:20], start=1):
        for col_idx, cell_val in enumerate(row_tuple):
            if cell_val is not None:
                s_val = str(cell_val).strip()
                if s_val == "FAI#ROW":
                    fai_col_idx = col_idx
                    start_data_row_idx = scan_idx
                    break
        if fai_col_idx is not None:
            break
    if fai_col_idx is None:
        raise Exception("在CPK工作表前20行未找到'FAI#ROW'表头，请检查文件！")

    # --------固定第9行（Excel物理行9，内存all_rows索引=8）--------
    row9_tuple = all_rows[8]
    val_row9 = row9_tuple[fai_col_idx] if fai_col_idx < len(row9_tuple) else None
    fai_text_row9 = str(val_row9) if val_row9 is not None else ""
    fixed_item = {
        "template_FAI#": fai_text_row9,
        "template_excel_row": 9,
        "title_column_letter": "A",
        "title_column_index": 1
    }

    fai_list = []
    # 表头下一行开始遍历全部行
    for mem_idx in range(start_data_row_idx, len(all_rows)):
        excel_physical_row = mem_idx + 1
        row_tuple = all_rows[mem_idx]
        val = row_tuple[fai_col_idx] if fai_col_idx < len(row_tuple) else None
        if val is None:
            fai_text = ""
        else:
            fai_text = str(val)
        # =========【改动点】：空内容直接跳过不加入列表 =========
        if fai_text.strip() == "":
            continue
        fai_list.append((fai_text, excel_physical_row))

    # ========= 2.只读title.csv表头 =========
    with open(title_path, "r", encoding="utf-8", newline="") as f:
        reader = csv.reader(f)
        title_header = next(reader)

    # =========3.逐行匹配 =========
    result_rows = []
    for fai_text, excel_row in fai_list:
        col_idx_1base = None
        if fai_text.strip() != "":
            try:
                col_idx_1base = title_header.index(fai_text) + 1
            except ValueError:
                col_idx_1base = None
        col_letter = col_num_to_letter(col_idx_1base)
        result_rows.append({
            "template_FAI#": fai_text,
            "template_excel_row": excel_row,
            "title_column_letter": col_letter,
            "title_column_index": col_idx_1base
        })
    full_sheet1 = [fixed_item] + result_rows

    # =========4.输出，仅Sheet1对应关系表 =========
    wb_out = Workbook()
    ws1 = wb_out.active
    ws1.title = "对应关系表"
    headers = ["template_FAI#", "template_excel_row", "title_column_letter", "title_column_index"]
    ws1.append(headers)
    for item in full_sheet1:
        ws1.append([
            item["template_FAI#"],
            item["template_excel_row"],
            item["title_column_letter"],
            item["title_column_index"]
        ])
    wb_out.save(output_path)
    return len(full_sheet1)

# ---------------- GUI界面 ----------------
class App(tk.Tk):
    def __init__(self, cfg):
        super().__init__()
        self.cfg = cfg
        self.title("FAI对应关系生成工具｜过滤空行｜FAI#ROW｜第9行固定A列")
        self.geometry("780x230")

        tk.Label(self, text="Template文件(xlsx)").grid(row=0, column=0, sticky="w", padx=8, pady=6)
        self.var_template = tk.StringVar(value=cfg.get("template_path",""))
        tk.Entry(self, textvariable=self.var_template, width=72).grid(row=0, column=1, padx=5)
        tk.Button(self, text="选择", command=self.select_template).grid(row=0, column=2, padx=4)

        tk.Label(self, text="Title文件(csv)").grid(row=1, column=0, sticky="w", padx=8, pady=6)
        self.var_title = tk.StringVar(value=cfg.get("title_path",""))
        tk.Entry(self, textvariable=self.var_title, width=72).grid(row=1, column=1, padx=5)
        tk.Button(self, text="选择", command=self.select_title).grid(row=1, column=2, padx=4)

        tk.Label(self, text="输出对应关系表(xlsx)").grid(row=2, column=0, sticky="w", padx=8, pady=6)
        self.var_output = tk.StringVar(value=cfg.get("output_path",""))
        tk.Entry(self, textvariable=self.var_output, width=72).grid(row=2, column=1, padx=5)
        tk.Button(self, text="选择", command=self.select_output).grid(row=2, column=2, padx=4)

        btn_gen = tk.Button(self, text="生成对应关系表", command=self.do_generate, bg="#4488dd", fg="white")
        btn_gen.grid(row=4, column=1, pady=20)

    def select_template(self):
        p = filedialog.askopenfilename(filetypes=[("Excel files","*.xlsx")])
        if p:
            self.var_template.set(p)

    def select_title(self):
        p = filedialog.askopenfilename(filetypes=[("CSV files","*.csv")])
        if p:
            self.var_title.set(p)

    def select_output(self):
        p = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel files","*.xlsx")])
        if p:
            self.var_output.set(p)

    def do_generate(self):
        tpl_path = self.var_template.get().strip()
        title_path = self.var_title.get().strip()
        out_path = self.var_output.get().strip()
        if not all([tpl_path, title_path, out_path]):
            messagebox.showerror("错误","请全部选择三个路径！")
            return
        if not os.path.exists(tpl_path):
            messagebox.showerror("错误","Template文件不存在！")
            return
        if not os.path.exists(title_path):
            messagebox.showerror("错误","Title文件不存在！")
            return
        new_cfg = {
            "template_path": tpl_path,
            "title_path": title_path,
            "output_path": out_path
        }
        save_config(new_cfg)
        try:
            total = run_mapping(tpl_path, title_path, out_path)
            messagebox.showinfo("完成", f"处理完成！共输出 {total} 行记录（包含第9行固定A列映射，已过滤空FAI#ROW行）\n输出文件：{out_path}")
        except Exception as e:
            messagebox.showerror("程序异常", f"执行失败：{str(e)}")

def main():
    cfg = load_config()
    # 无界面模式
    if len(sys.argv) ==4 and sys.argv[1]=="--headless":
        tpl = sys.argv[2]
        tit = sys.argv[3]
        out = sys.argv[4]
        total = run_mapping(tpl, tit, out)
        print(f"headless模式执行完毕，输出{total}行（包含第9行固定A列映射，已过滤空FAI#ROW行）")
        return
    app = App(cfg)
    app.mainloop()

if __name__ == "__main__":
    main()